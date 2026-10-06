"""
Bộ đọc nguồn danh sách từ kiểu cắm thêm (plugin). Mỗi nguồn một module trong thư mục này, khai báo một lớp con của
`SourceReader` và đăng ký bằng `@register`. Thêm nguồn mới (vd. NGSL, TSL): ghi giấy phép vào docs/data-sources.md, viết
module mới, import nó trong `_load_plugins` — KHÔNG phải sửa các bước 01–07.

Mỗi bộ đọc trả `RawRow` thô (headword, từ loại, cấp CEFR, thứ hạng nếu nguồn có); chuẩn hóa nằm ở lib/normalize.py.
"""

import csv
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class RawRow:
    headword: str
    pos: str | None
    cefr: str | None
    rank: int | None  # thứ hạng tần suất / thứ tự trong nguồn (nhỏ = phổ biến / dễ hơn); None nếu nguồn không có
    source: str
    extra: dict = field(default_factory=dict)


class SourceReader:
    name: str = ""
    has_frequency_rank: bool = False  # True: `rank` là thứ hạng tần suất thật; False: chỉ là thứ tự dòng

    def matches(self, path: Path) -> bool:
        raise NotImplementedError

    def version(self, path: Path) -> str | None:
        return None

    def read(self, path: Path) -> list[RawRow]:
        raise NotImplementedError


REGISTRY: list[SourceReader] = []


def register(cls):
    REGISTRY.append(cls())
    return cls


def _load_plugins() -> None:
    from data_pipeline.lib.sources import cefrj  # noqa: F401


def reader_for(path: Path) -> SourceReader | None:
    if not REGISTRY:
        _load_plugins()
    return next((r for r in REGISTRY if r.matches(path)), None)


def readers() -> list[SourceReader]:
    if not REGISTRY:
        _load_plugins()
    return list(REGISTRY)


def read_table(path: Path, required: set[str]) -> list[dict[str, str]]:
    """Đọc CSV / TSV / XLSX thành list dict (tên cột chữ thường, cắt khoảng trắng). Với XLSX: lấy sheet đầu tiên có đủ
    các cột `required` ở một trong 10 dòng đầu (dòng tiêu đề có thể không nằm ở dòng 1)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        text = path.read_text(encoding="utf-8-sig")
        delimiter = "\t" if suffix == ".tsv" or text.split("\n", 1)[0].count("\t") > text.split("\n", 1)[0].count(",") else ","
        rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
        return _rows_to_dicts(rows, required, path)
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
                try:
                    return _rows_to_dicts(rows, required, path)
                except ValueError:
                    continue
        finally:
            wb.close()
        raise ValueError(f"{path.name}: không sheet nào có đủ cột {sorted(required)}")
    raise ValueError(f"{path.name}: định dạng chưa hỗ trợ ({suffix})")


def _rows_to_dicts(rows: list[list[str]], required: set[str], path: Path) -> list[dict[str, str]]:
    for i, row in enumerate(rows[:10]):
        header = [str(c).strip().lower() for c in row]
        if required <= set(header):
            out = []
            for r in rows[i + 1:]:
                if not any(str(c).strip() for c in r):
                    continue
                out.append({header[j]: str(r[j]).strip() if j < len(r) else "" for j in range(len(header)) if header[j]})
            return out
    raise ValueError(f"{path.name}: thiếu cột {sorted(required)}")
